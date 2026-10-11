import calendar
import json
import secrets
import string
from datetime import date as date_cls, timedelta

from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, get_object_or_404, redirect
from django.templatetags.static import static
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .models import Classroom, Parent, Student, Attendance, Grade


def is_parent(user):
    return hasattr(user, "parent")


def classrooms_for(user):
    """المدير يشوف كل الفصول، والمدرس فصوله بس."""
    if user.is_staff:
        return Classroom.objects.all()
    return Classroom.objects.filter(teacher=user)


@login_required
def home(request):
    if is_parent(request.user):
        return redirect("parent_dashboard")
    if request.user.is_staff:
        return redirect("manager_dashboard")
    return redirect("classroom_list")


@login_required
def classroom_list(request):
    if is_parent(request.user):
        return redirect("parent_dashboard")
    return render(request, "school/classroom_list.html",
                  {"classrooms": classrooms_for(request.user)})


@login_required
def take_attendance(request, pk):
    if is_parent(request.user):
        return redirect("parent_dashboard")
    classroom = get_object_or_404(classrooms_for(request.user), pk=pk)
    students = list(classroom.student_set.all())

    raw = request.POST.get("date") if request.method == "POST" else request.GET.get("date")
    try:
        day = date_cls.fromisoformat(raw) if raw else timezone.localdate()
    except ValueError:
        messages.error(request, "تاريخ غير صحيح")
        day = timezone.localdate()

    if request.method == "POST":
        for s in students:
            status = request.POST.get(f"status_{s.id}")
            if status in ("P", "A", "L"):
                Attendance.objects.update_or_create(
                    student=s, date=day, defaults={"status": status}
                )
        messages.success(request, "تم حفظ الحضور")
        return redirect(f"{request.path}?date={day.isoformat()}")

    existing = {
        a.student_id: a.status
        for a in Attendance.objects.filter(student__in=students, date=day)
    }
    rows = [{"student": s, "status": existing.get(s.id, "P")} for s in students]

    return render(request, "school/take_attendance.html", {
        "classroom": classroom,
        "rows": rows,
        "today": day.isoformat(),
        "already_saved": bool(existing),
    })


@login_required
def add_grades(request, pk):
    if is_parent(request.user):
        return redirect("parent_dashboard")
    classroom = get_object_or_404(classrooms_for(request.user), pk=pk)
    students = classroom.student_set.all()

    if request.method == "POST":
        subject = request.POST.get("subject", "").strip()
        max_score = request.POST.get("max_score") or 100
        if not subject:
            messages.error(request, "اكتبي اسم المادة")
        else:
            for s in students:
                score = request.POST.get(f"score_{s.id}")
                if score:
                    Grade.objects.create(student=s, subject=subject,
                                         score=score, max_score=max_score)
            messages.success(request, "تم حفظ الدرجات")
        return redirect("add_grades", pk=pk)

    return render(request, "school/add_grades.html",
                  {"classroom": classroom, "students": students})


@login_required
def parent_dashboard(request):
    parent = getattr(request.user, "parent", None)
    if parent is None:
        return redirect("classroom_list")

    children = []
    for child in parent.children.all():
        records = child.attendance_set.order_by("-date")
        total = records.count()
        attended = records.filter(status__in=["P", "L"]).count()

        grades = child.grade_set.order_by("-date", "-id")
        pcts = [g.percent() for g in grades]
        avg = round(sum(pcts) / len(pcts)) if pcts else None
        if avg is None:
            level = None
        elif avg >= 85:
            level = "ممتاز"
        elif avg >= 65:
            level = "جيد"
        else:
            level = "يحتاج متابعة"

        latest = list(grades[:8])[::-1]
        chart_json = json.dumps({
            "labels": [g.subject for g in latest],
            "scores": [g.percent() for g in latest],
            "att": [
                records.filter(status="P").count(),
                records.filter(status="A").count(),
                records.filter(status="L").count(),
            ],
        }, ensure_ascii=False)

        children.append({
            "student": child,
            "records": records[:10],
            "total": total,
            "absent": records.filter(status="A").count(),
            "percent": round(attended * 100 / total) if total else None,
            "grades": grades[:10],
            "avg": avg,
            "level": level,
            "chart_json": chart_json,
        })
    return render(request, "school/parent_dashboard.html", {"children": children})


@login_required
def export_attendance(request, pk):
    if not request.user.is_staff:          # المدير بس
        return redirect("home")
    classroom = get_object_or_404(Classroom, pk=pk)

    today = timezone.localdate()
    try:
        year = int(request.GET.get("year", today.year))
        month = int(request.GET.get("month", today.month))
        days = calendar.monthrange(year, month)[1]
    except ValueError:
        return redirect("take_attendance", pk=pk)

    records = Attendance.objects.filter(
        student__classroom=classroom, date__year=year, date__month=month
    )
    by_day = {(r.student_id, r.date.day): r.status for r in records}
    mark = {"P": "ح", "A": "غ", "L": "ت"}

    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d}"
    ws.sheet_view.rightToLeft = True

    ws.append([f"حضور فصل {classroom.name} - {month}/{year}"])
    ws.append(["الطالب"] + list(range(1, days + 1)) + ["حاضر", "غائب", "متأخر", "نسبة الحضور"])

    for s in classroom.student_set.all():
        statuses = [by_day.get((s.id, d)) for d in range(1, days + 1)]
        p, a, l = statuses.count("P"), statuses.count("A"), statuses.count("L")
        total = p + a + l
        pct = (p + l) / total if total else None
        ws.append([s.name] + [mark.get(x, "") for x in statuses] + [p, a, l, pct])
        if pct is not None:
            ws.cell(row=ws.max_row, column=days + 5).number_format = "0%"

    ws["A1"].font = Font(bold=True, size=14)
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 24
    for row in ws.iter_rows(min_row=3, min_col=2):
        for c in row:
            c.alignment = Alignment(horizontal="center")

    resp = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp["Content-Disposition"] = (
        f'attachment; filename="attendance_{classroom.id}_{year}-{month:02d}.xlsx"'
    )
    wb.save(resp)
    return resp


@login_required
def manager_dashboard(request):
    if not request.user.is_staff:
        return redirect("home")

    today = timezone.localdate()
    since = today - timedelta(days=30)

    # ملخص النهاردة
    todays = Attendance.objects.filter(date=today)
    p = todays.filter(status="P").count()
    a = todays.filter(status="A").count()
    l = todays.filter(status="L").count()
    total = p + a + l
    today_percent = round((p + l) * 100 / total) if total else None
    absent_today = (todays.filter(status="A")
                    .select_related("student", "student__classroom")
                    .order_by("student__classroom__name", "student__name"))

    # أكتر الطلاب غيابًا آخر 30 يوم
    top_absent = (Student.objects
                  .annotate(absences=Count(
                      "attendance",
                      filter=Q(attendance__status="A", attendance__date__gte=since)))
                  .filter(absences__gt=0)
                  .select_related("classroom")
                  .order_by("-absences")[:5])

    # نسبة حضور كل فصل آخر 30 يوم
    class_rows = []
    classes = Classroom.objects.annotate(
        students=Count("student", distinct=True),
        total=Count("student__attendance",
                    filter=Q(student__attendance__date__gte=since)),
        attended=Count("student__attendance",
                       filter=Q(student__attendance__date__gte=since,
                                student__attendance__status__in=["P", "L"])),
    ).select_related("teacher")
    for c in classes:
        class_rows.append({
            "classroom": c,
            "percent": round(c.attended * 100 / c.total) if c.total else None,
        })

    return render(request, "school/manager_dashboard.html", {
        "today": today,
        "total_students": Student.objects.count(),
        "total_classes": Classroom.objects.count(),
        "p": p, "a": a, "l": l, "total": total,
        "today_percent": today_percent,
        "absent_today": absent_today,
        "top_absent": top_absent,
        "class_rows": class_rows,
    })


def _cell(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip() if v is not None else ""


@login_required
def import_students(request):
    if not request.user.is_staff:
        return redirect("home")

    results = []
    if request.method == "POST" and request.FILES.get("file"):
        try:
            ws = load_workbook(request.FILES["file"], read_only=True).active
        except Exception:
            messages.error(request, "الملف مش Excel صالح (لازم .xlsx)")
            return redirect("import_students")

        alphabet = string.ascii_letters + string.digits
        added = skipped = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            cells = [_cell(c) for c in (list(row) + [None] * 4)[:4]]
            name, cname, pname, phone = cells
            if not any(cells):
                continue
            if not name or not cname:
                skipped += 1
                continue

            classroom, _ = Classroom.objects.get_or_create(name=cname)
            parent = None
            if phone:
                user = User.objects.filter(username=f"p{phone}").first()
                if user is None:
                    pwd = "".join(secrets.choice(alphabet) for _ in range(8))
                    user = User.objects.create_user(
                        username=f"p{phone}", password=pwd, first_name=pname)
                    results.append({"parent": pname or phone,
                                    "username": user.username, "password": pwd})
                parent, _ = Parent.objects.get_or_create(
                    user=user, defaults={"phone": phone})

            _, new = Student.objects.get_or_create(
                name=name, classroom=classroom, defaults={"parent": parent})
            added += int(new)

        messages.success(request, f"تم إضافة {added} طالب، وتخطي {skipped} صف ناقص")

    return render(request, "school/import_students.html", {"results": results})


class TeacherForm(forms.Form):
    full_name = forms.CharField(label="اسم المدرس", max_length=150)
    username = forms.CharField(label="اسم المستخدم", max_length=150)
    password = forms.CharField(label="كلمة المرور", widget=forms.PasswordInput)
    classrooms = forms.ModelMultipleChoiceField(
        queryset=Classroom.objects.all(), required=False,
        label="الفصول", widget=forms.CheckboxSelectMultiple)

    def clean_username(self):
        u = self.cleaned_data["username"].strip()
        if User.objects.filter(username=u).exists():
            raise ValidationError("اسم المستخدم موجود بالفعل")
        return u

    def clean(self):
        data = super().clean()
        pwd = data.get("password")
        if pwd:
            try:
                validate_password(pwd)
            except ValidationError as e:
                self.add_error("password", e)
        return data


@login_required
def manage_teachers(request):
    if not request.user.is_staff:
        return redirect("home")

    form = TeacherForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        d = form.cleaned_data
        user = User.objects.create_user(
            username=d["username"], password=d["password"], first_name=d["full_name"])
        for c in d["classrooms"]:
            c.teacher = user
            c.save()
        messages.success(request, f"تم إضافة المدرس {d['full_name']}")
        return redirect("manage_teachers")

    teachers = (User.objects.filter(is_staff=False, parent__isnull=True)
                .prefetch_related("classrooms").order_by("first_name"))
    return render(request, "school/manage_teachers.html",
                  {"form": form, "teachers": teachers})


# ---------- تطبيق الموبايل (PWA): من غير تسجيل دخول ----------

def manifest(request):
    return JsonResponse({
        "name": "منصة المتابعة الدراسية",
        "short_name": "المتابعة",
        "start_url": "/",
        "scope": "/",
        "display": "standalone",
        "background_color": "#f3f6fb",
        "theme_color": "#1e3a8a",
        "lang": "ar",
        "dir": "rtl",
        "icons": [
            {"src": static("school/icons/icon-192.png"), "sizes": "192x192", "type": "image/png"},
            {"src": static("school/icons/icon-512.png"), "sizes": "512x512", "type": "image/png"},
            {"src": static("school/icons/icon-512.png"), "sizes": "512x512",
             "type": "image/png", "purpose": "maskable"},
        ],
    }, json_dumps_params={"ensure_ascii": False})


def service_worker(request):
    js = """
self.addEventListener('install', function (e) { self.skipWaiting(); });
self.addEventListener('activate', function (e) { e.waitUntil(self.clients.claim()); });
self.addEventListener('fetch', function (e) {
  if (e.request.mode === 'navigate') {
    e.respondWith(fetch(e.request).catch(function () {
      return new Response(
        '<html lang="ar" dir="rtl"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">' +
        '<body style="font-family:sans-serif;text-align:center;padding:40px">' +
        '<h2>مفيش اتصال بالإنترنت</h2><p>اتأكد من الاتصال وحاول تاني.</p></body></html>',
        { headers: { 'Content-Type': 'text/html; charset=utf-8' } });
    }));
  }
});
"""
    return HttpResponse(js, content_type="application/javascript")